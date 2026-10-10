#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Planilha do bloco de força e potência da semana do jogo (12 a 16/10/2026).

Uma entrada só: o 1RM de cada atleta na aba 1RM. Todo percentual das quatro
sessões vira quilo por fórmula, arredondado de 2,5 em 2,5 kg — que é o que a
anilha permite. Mudou o 1RM, mudou a sessão inteira.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONTE = "Arial"
TINTA   = "1A1A1A"
TINTA_2 = "5A6672"
AZUL    = "2E6DB4"
AZUL_CL = "E8F0FA"
CINZA   = "F2F2F2"
AMARELO = "FFF2A8"      # célula que o usuário preenche
LINHA   = "BFBFBF"
LINHA_L = "D9D9D9"

fino  = Side(style="thin", color=LINHA_L)
medio = Side(style="thin", color=LINHA)
BORDA      = Border(left=fino, right=fino, top=fino, bottom=fino)
BORDA_CAB  = Border(left=fino, right=fino, top=medio, bottom=medio)

# ── Os exercícios de referência: a MESMA lista do app da comissão ──────────
# A ordem e a grafia são protocolo: a fórmula de cada sessão procura o nome
# nesta coluna. Um nome diferente aqui é uma carga que não aparece.
REFS = ["Agachamento", "Agachamento frontal", "Supino", "Stiff com barra",
        "Remada serrote", "Arranco", "Clean", "Clean pull", "Push press"]
EXEMPLO = {"Agachamento":150, "Agachamento frontal":125, "Supino":110,
           "Stiff com barra":120, "Remada serrote":80, "Arranco":70,
           "Clean":95, "Clean pull":110, "Push press":75}

N_ATLETAS = 14          # elenco de voleibol cabe; sobra coluna é só deixar vazia

# (nº, exercício, padrão, séries, reps, pausa_s, carga, ref1rm, pct, observação)
SEGUNDA = [
 (1,"Agachamento","Perna",5,3,180,"85%","Agachamento",0.85,
  "Âncora do bloco. Se a terceira repetição sair lenta, a série acabou."),
 (2,"Levantamento terra","Perna",4,3,180,"RIR 1",None,None,
  "Sem referência de 1RM no sistema: vai por repetições na reserva. Uma de sobra, nunca zero."),
 (3,"Avanço com barra, à frente e atrás","Perna",3,5,150,"RIR 2",None,None,
  "Cinco ciclos em CADA perna — um passo à frente e um atrás contam como um."),
 (4,"Stiff com barra","Perna",4,5,120,"72%","Stiff com barra",0.72,
  "Excêntrica de 3 s (3-0-1-0). A carga é menor de propósito: quem manda aqui é o tempo sob tensão."),
 (5,"Elevação pélvica unilateral","Perna",3,10,90,"RIR 2",None,None,
  "Dez em cada perna."),
 (6,"Nórdico de isquiotibiais (excêntrico)","Perna",3,6,90,"peso do corpo",None,None,
  "É a dose que reduz lesão de posterior. Descida lenta até perder o controle."),
 (7,"Elevação de calcanhares","Perna",3,8,60,"RIR 2",None,None,
  "Descida de 3 s."),
 (8,"Abdominal reto com braços esticados","Core",3,20,45,"RIR 2",None,None,
  ""),
]
TERCA = [
 (1,"Supino reto com barra","Empurrar",5,3,180,"85%","Supino",0.85,
  "Mesma lógica da segunda: três repetições sólidas valem mais que quatro arrastadas."),
 (2,"Push press","Empurrar",4,4,150,"78%","Push press",0.78,
  "A perna dá o arranque, o braço só termina. Se o braço faz tudo, a carga está alta demais."),
 (3,"Remada serrote","Puxar",4,6,120,"78%","Remada serrote",0.78,
  "Seis em cada braço."),
 (4,"Puxada unilateral no pulley alto","Puxar",3,8,90,"RIR 2",None,None,
  "Oito em cada braço."),
 (5,"Remada alta","Puxar",3,8,90,"RIR 2",None,None,
  ""),
 (6,"Rotadores do ombro com elástico","Ombro",3,15,60,"RIR 3",None,None,
  "Força do estabilizador, não mobilidade. As duas coisas entram e são diferentes."),
 (7,"Y-T-W no banco inclinado","Ombro",2,8,45,"carga leve",None,None,
  "Oito de cada letra. Movimento da escápula, não do braço."),
 (8,"Prancha lateral com elevação de quadril","Core",3,10,45,"peso do corpo",None,None,
  "Dez de cada lado."),
]
QUARTA = [
 (1,"Clean","LPO",5,2,180,"80%","Clean",0.80,
  "Velocidade da barra manda. Se cair, a série acabou — mesmo que sobre repetição."),
 (2,"Snatch pull / Hang high pull","LPO",3,3,150,"90%","Clean pull",0.90,
  "Puxada alta: a parte do arranco que aguenta mais carga."),
 (3,"Agachamento com salto sob carga (jump squat)","Potência",4,4,180,"25%","Agachamento",0.25,
  "Acima de 1,0 m/s. Carga leve de propósito: aqui se treina velocidade, não força."),
 (4,"Supino com a barra partindo do peito","Potência",4,3,150,"55%","Supino",0.55,
  "A barra sai PARADA do peito e vai na maior velocidade possível. Sem quique."),
 (5,"Drop jump 40 cm","Pliometria",3,5,120,"peso do corpo",None,None,
  "Contato curto com o solo. Qualidade acima da contagem."),
 (6,"Salto no caixote","Pliometria",3,5,90,"peso do corpo",None,None,
  "Subir saltando, DESCER andando: descer do caixote é contato extra sem ganho."),
 (7,"Sprints de 10 e 20 m com mudança de direção","Potência",4,1,120,"peso do corpo",None,None,
  "Recuperação completa entre tiros."),
]
QUINTA = [
 (1,"Clean do joelho (hang)","LPO",3,2,120,"70%","Clean",0.70,
  "Técnico e rápido. Nenhuma tentativa pesada na véspera."),
 (2,"Agachamento com salto sob carga (jump squat)","Potência",3,3,120,"20%","Agachamento",0.20,
  "Mais leve que na quarta. O objetivo é acordar o sistema, não cansá-lo."),
 (3,"Salto no caixote","Pliometria",3,3,90,"peso do corpo",None,None,
  "Altura confortável. Subir saltando, descer andando."),
 (4,"Arremesso de medicine ball acima da cabeça","Potência",3,4,60,"3 a 4 kg",None,None,
  "Lançamento à frente, com os dois braços, o mais longe possível."),
 (5,"Mobilização de tornozelo na parede (knee-to-wall)","Mobilidade",2,8,30,"peso do corpo",None,None,
  "Oito em cada lado, calcanhar no chão. Abre o agachamento e a aterrissagem."),
]

DIAS = [
 ("Segunda", "SEGUNDA 12/10", "FORÇA 1 · dominante de perna",
  "A sessão mais pesada da semana, e a mais longe do jogo: 96 horas para a "
  "fadiga sair das pernas.", SEGUNDA),
 ("Terça", "TERÇA 13/10", "FORÇA 2 · dominante de tronco",
  "Carga alta de novo, mas em outro tecido: a perna de ontem descansa "
  "enquanto o tronco trabalha.", TERCA),
 ("Quarta", "QUARTA 14/10", "POTÊNCIA 1 · completa",
  "A virada do bloco. Percentual cai, velocidade manda — e entram os "
  "contatos pliométricos da semana.", QUARTA),
 ("Quinta", "QUINTA 15/10", "POTÊNCIA 2 · véspera do jogo",
  "Alta intensidade e volume BAIXO: 24 minutos. É para acordar o sistema "
  "nervoso, não para treinar.", QUINTA),
]

LIN_PRIM = 8            # primeira linha de exercício em cada aba do dia
COL_AT_1 = 11           # coluna K: primeiro atleta
COL_RM_1 = 2            # coluna B na aba 1RM: primeiro atleta
RM_LIN_1, RM_LIN_N = 4, 4 + len(REFS) - 1

wb = Workbook()


def titulo(ws, linha, texto, tam=14, cor=TINTA, negrito=True):
    c = ws.cell(row=linha, column=1, value=texto)
    c.font = Font(name=FONTE, size=tam, bold=negrito, color=cor)
    return c


def cabecalho(ws, linha, rotulos, larguras):
    for i, (rot, larg) in enumerate(zip(rotulos, larguras), start=1):
        c = ws.cell(row=linha, column=i, value=rot)
        c.font = Font(name=FONTE, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA_CAB
        if larg:
            ws.column_dimensions[get_column_letter(i)].width = larg


# ══════════════════════════════════════════════════════════════════════════
# Aba 1RM — a única entrada da planilha
# ══════════════════════════════════════════════════════════════════════════
rm = wb.active
rm.title = "1RM"
titulo(rm, 1, "1RM DO ELENCO", 15, AZUL)
c = rm.cell(row=2, column=1, value="A única coisa que você digita nesta planilha. "
            "Cada percentual das quatro sessões vira quilo a partir daqui, "
            "arredondado de 2,5 em 2,5 kg.")
c.font = Font(name=FONTE, size=9.5, color=TINTA_2)

rotulos = ["Exercício de referência", "EXEMPLO (apague)"] + \
          ["Atleta %d" % i for i in range(1, N_ATLETAS + 1)]
larguras = [30, 15] + [11] * N_ATLETAS
cabecalho(rm, 3, rotulos, larguras)

for i, ref in enumerate(REFS):
    lin = RM_LIN_1 + i
    c = rm.cell(row=lin, column=1, value=ref)
    c.font = Font(name=FONTE, size=10, bold=True)
    c.border = BORDA
    c.alignment = Alignment(vertical="center")
    for j in range(N_ATLETAS + 1):            # +1 por causa da coluna EXEMPLO
        col = COL_RM_1 + j
        cel = rm.cell(row=lin, column=col)
        if j == 0:
            cel.value = EXEMPLO[ref]
        cel.fill = PatternFill("solid", fgColor=AMARELO)
        cel.border = BORDA
        cel.number_format = "0.0"
        cel.alignment = Alignment(horizontal="center")
        cel.font = Font(name=FONTE, size=10, color="0000FF")

lin = RM_LIN_N + 2
for texto, negrito in [
    ("Como usar", True),
    ("1.  Troque \"Atleta 1\", \"Atleta 2\"… pelos nomes do elenco, na linha 3.", False),
    ("2.  Digite o 1RM de cada um nas células amarelas. Só o que foi testado.", False),
    ("3.  Deixe EM BRANCO o que não foi medido — a sessão mostra em branco também,", False),
    ("     e campo vazio é melhor que número chutado: é dele que sai a carga do bloco inteiro.", False),
    ("4.  A coluna EXEMPLO existe só para você ver o formato funcionando. Apague-a.", False),
    ("", False),
    ("Estes nove exercícios são os mesmos do app da comissão, com a mesma grafia. "
     "Mudar um nome aqui faz a carga sumir das sessões.", False),
]:
    c = rm.cell(row=lin, column=1, value=texto)
    c.font = Font(name=FONTE, size=9.5, bold=negrito,
                  color=TINTA if negrito else TINTA_2)
    lin += 1

rm.freeze_panes = "B4"


# ══════════════════════════════════════════════════════════════════════════
# Uma aba por dia
# ══════════════════════════════════════════════════════════════════════════
def aba_dia(nome, dia, foco, chamada, exs):
    ws = wb.create_sheet(nome)
    titulo(ws, 1, "ELASE VOLEIBOL · MASCULINO ADULTO · SEMANA DO JOGO", 9, TINTA_2, False)
    titulo(ws, 2, "%s  ·  %s" % (dia, foco), 15, AZUL)
    c = ws.cell(row=3, column=1, value=chamada)
    c.font = Font(name=FONTE, size=10, color=TINTA_2)

    ult = LIN_PRIM + len(exs) - 1
    resumo = [
        ("Séries", "=SUM(D%d:D%d)" % (LIN_PRIM, ult), "0"),
        ("Duração estimada",
         "=ROUND(SUMPRODUCT($D$%d:$D$%d,($F$%d:$F$%d+$E$%d:$E$%d*4))/60,0)"
         % (LIN_PRIM, ult, LIN_PRIM, ult, LIN_PRIM, ult), '0" min"'),
        ("Contatos pliométricos",
         '=SUMPRODUCT(--($C$%d:$C$%d="Pliometria"),$D$%d:$D$%d,$E$%d:$E$%d)'
         % (LIN_PRIM, ult, LIN_PRIM, ult, LIN_PRIM, ult), "0"),
    ]
    for i, (rot, form, fmt) in enumerate(resumo):
        col = 1 + i * 2
        a = ws.cell(row=5, column=col, value=rot)
        a.font = Font(name=FONTE, size=8.5, bold=True, color=TINTA_2)
        b = ws.cell(row=6, column=col, value=form)
        b.font = Font(name=FONTE, size=13, bold=True, color=AZUL)
        b.number_format = fmt

    rot = ["#", "Exercício", "Padrão", "Séries", "Reps", "Pausa (s)", "Carga",
           "Ref. de 1RM", "%", "Observação"]
    larg = [4, 38, 11, 7, 6, 9, 11, 17, 7, 46]
    nomes_at = []
    for j in range(N_ATLETAS + 1):
        rot.append("")
        larg.append(10)
    cabecalho(ws, 7, rot, larg)
    # O nome do atleta vem da aba 1RM: trocar lá troca nas quatro sessões.
    for j in range(N_ATLETAS + 1):
        col = COL_AT_1 + j
        letra_rm = get_column_letter(COL_RM_1 + j)
        c = ws.cell(row=7, column=col,
                    value="=IF('1RM'!%s$3=\"\",\"\",'1RM'!%s$3)" % (letra_rm, letra_rm))
        c.font = Font(name=FONTE, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA_CAB
        nomes_at.append(col)

    for i, (n, nome_ex, padrao, series, reps, pausa, carga, ref, pct, obs) in enumerate(exs):
        lin = LIN_PRIM + i
        zebra = PatternFill("solid", fgColor=CINZA) if i % 2 else None
        valores = [n, nome_ex, padrao, series, reps, pausa, carga, ref or "—",
                   pct if pct is not None else "", obs]
        for k, v in enumerate(valores, start=1):
            c = ws.cell(row=lin, column=k, value=v)
            c.border = BORDA
            if zebra:
                c.fill = zebra
            c.font = Font(name=FONTE, size=10,
                          bold=(k == 2), color=TINTA if k != 10 else TINTA_2)
            if k == 10:
                c.font = Font(name=FONTE, size=8.5, color=TINTA_2)
                c.alignment = Alignment(vertical="top", wrap_text=True)
            elif k in (1, 3, 4, 5, 6, 7, 8, 9):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(vertical="center", wrap_text=True)
            if k == 9 and pct is not None:
                c.number_format = "0%"
        # Altura pela observação, que é o texto mais longo da linha. Fixa em 32,
        # a de 95 caracteres ficava cortada na coluna de 46 — e observação
        # cortada numa ficha de treino é instrução que ninguém lê.
        linhas_obs = max(2, -(-len(obs) // 42), -(-len(nome_ex) // 36))
        ws.row_dimensions[lin].height = 10 + linhas_obs * 12

        for j, col in enumerate(nomes_at):
            letra_rm = get_column_letter(COL_RM_1 + j)
            faixa = "'1RM'!%s$%d:%s$%d" % (letra_rm, RM_LIN_1, letra_rm, RM_LIN_N)
            achar = "MATCH($H%d,'1RM'!$A$%d:$A$%d,0)" % (lin, RM_LIN_1, RM_LIN_N)
            busca = "INDEX(%s,%s)" % (faixa, achar)
            # Sem percentual não há quilo a calcular; sem 1RM lançado a célula
            # fica VAZIA em vez de zero, que se leria como "pode usar zero".
            form = ('=IF($I{l}="","",IFERROR(IF({b}="","",'
                    'ROUND({b}*$I{l}/2.5,0)*2.5),""))').format(l=lin, b=busca)
            c = ws.cell(row=lin, column=col, value=form)
            c.border = BORDA
            if zebra:
                c.fill = zebra
            c.font = Font(name=FONTE, size=11, bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.number_format = "0.0"

    lin = ult + 2
    for texto, negrito in [
        ("Regra da sessão", True),
        ("A coluna de cada atleta é o quilo que sai do 1RM dele, já arredondado "
         "à anilha. Em branco = falta o 1RM daquele exercício na aba 1RM.", False),
        ("\"RIR 2\" é repetições na reserva: terminar a série com duas que ainda "
         "sairiam com técnica limpa. Não é ir até a falha.", False),
    ]:
        c = ws.cell(row=lin, column=1, value=texto)
        c.font = Font(name=FONTE, size=9.5, bold=negrito,
                      color=TINTA if negrito else TINTA_2)
        lin += 1

    ws.freeze_panes = ws.cell(row=LIN_PRIM, column=COL_AT_1)
    ws.sheet_view.showGridLines = False
    return ws


for nome, dia, foco, chamada, exs in DIAS:
    aba_dia(nome, dia, foco, chamada, exs)


# ══════════════════════════════════════════════════════════════════════════
# Aba Bloco — o mapa da semana e a contabilidade
# ══════════════════════════════════════════════════════════════════════════
bl = wb.create_sheet("Bloco", 0)
titulo(bl, 1, "BLOCO DE FORÇA E POTÊNCIA · SEMANA DO JOGO", 16, AZUL)
c = bl.cell(row=2, column=1, value="ELASE Voleibol · masculino adulto · 12 a 16 de outubro de 2026")
c.font = Font(name=FONTE, size=10, color=TINTA_2)

c = bl.cell(row=4, column=1, value=
    "Dois dias de muita força e dois de muita potência, nesta ordem, porque é a "
    "ordem que o jogo de sexta permite: a força pesada fica longe do jogo, onde "
    "há tempo para a fadiga sair, e a potência fica perto, onde ela ajuda em vez "
    "de atrapalhar.")
c.font = Font(name=FONTE, size=10, color=TINTA)
c.alignment = Alignment(wrap_text=True, vertical="top")
bl.merge_cells("A4:G5")
bl.row_dimensions[4].height = 22
bl.row_dimensions[5].height = 22

cabecalho(bl, 7, ["Dia", "Data", "Foco", "Séries", "Duração", "Contatos", "Por quê"],
          [11, 10, 30, 8, 10, 10, 62])

mapa = [
 ("Segunda", "12/10", "FORÇA 1 · perna",
  "A mais pesada e a mais longe do jogo: 96 h para a fadiga sair da perna."),
 ("Terça", "13/10", "FORÇA 2 · tronco",
  "Carga alta de novo, em outro tecido. A perna de ontem descansa enquanto o tronco trabalha."),
 ("Quarta", "14/10", "POTÊNCIA 1 · completa",
  "Percentual cai, velocidade manda. É aqui que entram os contatos pliométricos da semana."),
 ("Quinta", "15/10", "POTÊNCIA 2 · véspera",
  "Intensidade alta, volume baixo, 24 min. Acordar o sistema nervoso — não treinar."),
]
for i, (dia, data, foco, porque) in enumerate(mapa):
    lin = 8 + i
    vals = [dia, data, foco,
            "='%s'!A6" % dia, "='%s'!C6" % dia, "='%s'!E6" % dia, porque]
    for k, v in enumerate(vals, start=1):
        c = bl.cell(row=lin, column=k, value=v)
        c.border = BORDA
        c.font = Font(name=FONTE, size=10, bold=(k in (1, 3)))
        if k in (4, 5, 6):
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.number_format = '0" min"' if k == 5 else "0"
        elif k == 7:
            c.font = Font(name=FONTE, size=8.5, color=TINTA_2)
            c.alignment = Alignment(vertical="center", wrap_text=True)
        else:
            c.alignment = Alignment(vertical="center")
    bl.row_dimensions[lin].height = 30

lin = 12
vals = ["SEXTA", "16/10", "JOGO", "—", "—", "—",
        "Nada de sala. O bloco inteiro existe para chegar aqui."]
for k, v in enumerate(vals, start=1):
    c = bl.cell(row=lin, column=k, value=v)
    c.border = BORDA
    c.fill = PatternFill("solid", fgColor=AZUL_CL)
    c.font = Font(name=FONTE, size=10, bold=True, color=AZUL)
    c.alignment = Alignment(horizontal="center" if k in (4, 5, 6) else "left",
                            vertical="center", wrap_text=(k == 7))
bl.row_dimensions[12].height = 24

lin = 13
for k, rot in enumerate(["TOTAL", "", "4 sessões",
                         "=SUM(D8:D11)", "=SUM(E8:E11)", "=SUM(F8:F11)",
                         "Mais o que a quadra cobrar nos mesmos quatro dias."],
                        start=1):
    c = bl.cell(row=lin, column=k, value=rot)
    c.border = BORDA
    c.font = Font(name=FONTE, size=10, bold=True)
    if k in (4, 5, 6):
        c.alignment = Alignment(horizontal="center")
        c.number_format = '0" min"' if k == 5 else "0"
    elif k == 7:
        c.font = Font(name=FONTE, size=8.5, color=TINTA_2)

lin = 15
blocos_texto = [
    ("O que decide este bloco", True, None),
    ("A ordem não é preferência, é prazo. Força a 85% deixa fadiga residual por "
     "48 a 72 horas; potência em volume baixo quase não deixa, e ainda potencia. "
     "Por isso a força vai para segunda e terça, e a potência para quarta e quinta.", False, None),
    ("As duas sessões de força não repetem o tecido: segunda é perna, terça é "
     "tronco. Duas sessões pesadas seguidas no mesmo grupo não seriam duas "
     "sessões — seria uma boa e uma arrastada.", False, None),
    ("", False, None),
    ("A quinta é a exceção que você precisa conhecer", True, None),
    ("É dia de potência DE VERDADE em intensidade — velocidade máxima, barra "
     "leve, salto alto — e volume deliberadamente baixo: 14 séries, 24 minutos, "
     "9 contatos. Véspera de jogo com volume de potência alto chega na sexta com "
     "a perna pesada, e aí o bloco inteiro trabalhou contra o jogo.", False, None),
    ("Nenhuma série até a falha, nenhuma excêntrica lenta, nenhuma carga nova na "
     "quinta. Se sobrar dúvida entre fazer e não fazer um exercício, não faça.", False, "crit"),
    ("", False, None),
    ("O que eu não sei daqui", True, None),
    ("Se há treino de quadra nos mesmos quatro dias — e quase certamente há. "
     "Se houver, segunda e terça a sala vem DEPOIS da quadra, ou com algumas "
     "horas de intervalo; e se a quarta tiver treino tático pesado, corte os "
     "sprints.", False, None),
    ("Os contatos pliométricos contados aqui são só os de pliometria pura "
     "(drop jump e salto no caixote), que é como o app da comissão conta. "
     "As aterrissagens do jump squat e tudo o que a quadra cobra ficam de fora "
     "da conta — some-os por cima antes de decidir que 39 é pouco.", False, None),
]
for texto, negrito, tom in blocos_texto:
    c = bl.cell(row=lin, column=1, value=texto)
    cor = TINTA if negrito else ("9B3B3B" if tom == "crit" else TINTA_2)
    c.font = Font(name=FONTE, size=10 if negrito else 9.5, bold=negrito, color=cor)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if texto and not negrito:
        bl.merge_cells(start_row=lin, start_column=1, end_row=lin, end_column=7)
        bl.row_dimensions[lin].height = 34
    lin += 1

lin += 1
for texto, negrito in [
    ("Como a planilha funciona", True),
    ("Abra a aba 1RM, ponha os nomes do elenco e o 1RM de cada um. As quatro "
     "abas de sessão calculam o quilo sozinhas, arredondado de 2,5 em 2,5 kg.", False),
    ("Célula amarela é para preencher. Todo o resto é fórmula — mexer nela "
     "quebra a conta.", False),
    ("1RM em branco deixa a carga em branco, de propósito: é melhor ver o vazio "
     "e perguntar do que receber um número que ninguém mediu.", False),
]:
    c = bl.cell(row=lin, column=1, value=texto)
    c.font = Font(name=FONTE, size=10 if negrito else 9.5, bold=negrito,
                  color=TINTA if negrito else TINTA_2)
    if texto and not negrito:
        bl.merge_cells(start_row=lin, start_column=1, end_row=lin, end_column=7)
    lin += 1

bl.sheet_view.showGridLines = False

for ws in wb:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

CAMINHO = "/home/user/mdlucca/docs/ELASE-bloco-forca-potencia-12a16-10.xlsx"
wb.save(CAMINHO)
print("ok:", CAMINHO)
